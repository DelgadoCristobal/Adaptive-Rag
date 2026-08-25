import datetime
from pathlib import Path
import yaml


class PPTXParser:
    """Parser for PPTX presentation slides using python-pptx."""

    def parse(self, file_path: Path) -> str:
        """Parse PPTX file into Obsidian-compatible Markdown with slide structure and notes."""
        file_path = Path(file_path)

        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            slide_sections = []

            for idx, slide in enumerate(prs.slides, start=1):
                slide_content = [f"## Slide {idx}"]
                
                # Extract text from shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        # Check if title shape
                        if shape == slide.shapes.title:
                            slide_content.insert(1, f"### {shape.text.strip()}")
                        else:
                            slide_content.append(shape.text.strip())

                # Speaker notes capture
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_content.append(f"\n> **Speaker Notes:** {notes}")

                slide_sections.append("\n\n".join(slide_content))

            markdown_content = f"# {file_path.stem}\n\n" + "\n\n---\n\n".join(slide_sections)
        except Exception as e:
            markdown_content = f"# {file_path.stem}\n\n*Error parsing PPTX: {e}*"

        metadata = {
            "title": file_path.stem,
            "source": file_path.name,
            "file_type": "pptx",
            "ingested_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "tags": ["presentation", "pptx"]
        }

        frontmatter = f"---\n{yaml.dump(metadata, sort_keys=False)}---\n\n"
        return frontmatter + markdown_content
