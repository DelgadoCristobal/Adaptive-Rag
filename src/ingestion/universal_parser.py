from typing import List, Dict, Any, Optional, Union
from enum import Enum
from pathlib import Path
import re
import logging
from pydantic import BaseModel, Field
from src.config import settings

logger = logging.getLogger(__name__)


class SourceType(str, Enum):
    PDF = "pdf"
    DOCX = "docx"
    PPTX = "pptx"
    MARKDOWN = "markdown"
    TEXT = "text"


class DocumentChunk(BaseModel):
    id: str
    doc_id: str
    source_type: SourceType
    page_or_slide_num: Optional[int] = None
    text: str
    metadata: Dict[str, Any] = Field(default_factory=dict)


class ProcessedDocument(BaseModel):
    doc_id: str
    filename: str
    source_type: SourceType
    title: str
    markdown_content: str
    chunks: List[DocumentChunk] = Field(default_factory=list)


class UniversalDocumentParser:
    """Universal Enterprise Document Ingestion Engine for PDF, DOCX, PPTX, MD, and TXT."""

    def __init__(self, markdown_dir: Optional[Path] = None, vault_dir: Optional[Path] = None):
        self.markdown_dir = markdown_dir or settings.data_markdown_dir
        self.vault_dir = vault_dir or settings.data_vault_dir
        self.markdown_dir.mkdir(parents=True, exist_ok=True)
        self.vault_dir.mkdir(parents=True, exist_ok=True)

    def parse_file(self, file_path: Union[str, Path]) -> ProcessedDocument:
        """Auto-detects file format and parses content into structured markdown."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self.parse_pdf(path)
        elif suffix in [".docx", ".doc"]:
            return self.parse_docx(path)
        elif suffix in [".pptx", ".ppt"]:
            return self.parse_pptx(path)
        elif suffix in [".md", ".markdown", ".txt"]:
            return self.parse_text_or_md(path)
        else:
            logger.warning(f"Unsupported extension '{suffix}'. Attempting plain text parsing.")
            return self.parse_text_or_md(path)

    def parse_pdf(self, file_path: Path) -> ProcessedDocument:
        """Extracts text, headers, and tables from PDF page-by-page using pymupdf4llm."""
        doc_id = self._clean_doc_id(file_path.stem)
        chunks: List[DocumentChunk] = []
        full_md_pages: List[str] = []

        try:
            import pymupdf4llm
            pages_data = pymupdf4llm.to_markdown(str(file_path), page_chunks=True)
            for idx, page in enumerate(pages_data, start=1):
                text = page.get("text", "").strip()
                if not text:
                    continue
                header = f"<!-- Page {idx} -->"
                full_md_pages.append(header + "\n" + text)
                chunks.append(DocumentChunk(
                    id=f"{doc_id}_p{idx}",
                    doc_id=doc_id,
                    source_type=SourceType.PDF,
                    page_or_slide_num=idx,
                    text=text,
                    metadata={"page": idx, "filename": file_path.name}
                ))
        except Exception as e:
            logger.error(f"Error parsing PDF {file_path}: {e}. Falling back to basic extraction.")
            import fitz
            doc = fitz.open(file_path)
            for idx, page in enumerate(doc, start=1):
                text = page.get_text().strip()
                if text:
                    header = f"<!-- Page {idx} -->"
                    full_md_pages.append(header + "\n" + text)
                    chunks.append(DocumentChunk(
                        id=f"{doc_id}_p{idx}",
                        doc_id=doc_id,
                        source_type=SourceType.PDF,
                        page_or_slide_num=idx,
                        text=text,
                        metadata={"page": idx, "filename": file_path.name}
                    ))

        full_markdown = "\n\n".join(full_md_pages)
        return ProcessedDocument(
            doc_id=doc_id,
            filename=file_path.name,
            source_type=SourceType.PDF,
            title=file_path.stem.replace("_", " ").title(),
            markdown_content=full_markdown,
            chunks=chunks
        )

    def parse_docx(self, file_path: Path) -> ProcessedDocument:
        """Extracts structured text, headings, and Markdown tables from DOCX."""
        import docx
        doc_id = self._clean_doc_id(file_path.stem)
        doc = docx.Document(file_path)

        md_elements: List[str] = []
        chunks: List[DocumentChunk] = []
        chunk_idx = 1

        for element in doc.element.body:
            if element.tag.endswith("p"):
                p = docx.text.paragraph.Paragraph(element, doc)
                text = p.text.strip()
                if not text:
                    continue
                style_name = p.style.name.lower() if p.style else ""
                if "heading 1" in style_name:
                    md_elements.append(f"# {text}")
                elif "heading 2" in style_name:
                    md_elements.append(f"## {text}")
                elif "heading 3" in style_name:
                    md_elements.append(f"### {text}")
                elif "list" in style_name or (p.style and p.style.name.startswith("List")):
                    md_elements.append(f"- {text}")
                else:
                    md_elements.append(text)

            elif element.tag.endswith("tbl"):
                table = docx.table.Table(element, doc)
                table_md = self._docx_table_to_markdown(table)
                if table_md:
                    md_elements.append(table_md)

        full_markdown = "\n\n".join(md_elements)

        sections = full_markdown.split("\n\n")
        current_chunk = []
        current_len = 0
        for sec in sections:
            current_chunk.append(sec)
            current_len += len(sec)
            if current_len >= 800:
                chunk_text = "\n\n".join(current_chunk).strip()
                chunks.append(DocumentChunk(
                    id=f"{doc_id}_c{chunk_idx}",
                    doc_id=doc_id,
                    source_type=SourceType.DOCX,
                    page_or_slide_num=None,
                    text=chunk_text,
                    metadata={"chunk_index": chunk_idx, "filename": file_path.name}
                ))
                chunk_idx += 1
                current_chunk = []
                current_len = 0

        if current_chunk:
            chunk_text = "\n\n".join(current_chunk).strip()
            chunks.append(DocumentChunk(
                id=f"{doc_id}_c{chunk_idx}",
                doc_id=doc_id,
                source_type=SourceType.DOCX,
                page_or_slide_num=None,
                text=chunk_text,
                metadata={"chunk_index": chunk_idx, "filename": file_path.name}
            ))

        return ProcessedDocument(
            doc_id=doc_id,
            filename=file_path.name,
            source_type=SourceType.DOCX,
            title=file_path.stem.replace("_", " ").title(),
            markdown_content=full_markdown,
            chunks=chunks
        )

    def parse_pptx(self, file_path: Path) -> ProcessedDocument:
        """Extracts slides, titles, bullet lists, Markdown tables, and speaker notes from PPTX."""
        from pptx import Presentation
        doc_id = self._clean_doc_id(file_path.stem)
        prs = Presentation(file_path)

        md_slides: List[str] = []
        chunks: List[DocumentChunk] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_lines: List[str] = [f"## Slide {slide_idx}"]
            
            if slide.shapes.title and slide.shapes.title.text:
                title_text = slide.shapes.title.text.strip()
                slide_lines.append(f"### {title_text}")

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        p_text = p.text.strip()
                        if p_text:
                            level = p.level
                            indent = "  " * level
                            slide_lines.append(f"{indent}- {p_text}")
                elif shape.has_table:
                    tbl_md = self._pptx_table_to_markdown(shape.table)
                    if tbl_md:
                        slide_lines.append(tbl_md)

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_lines.append(f"\n> **Speaker Notes (Slide {slide_idx}):** {notes_text}")

            slide_md = "\n\n".join(slide_lines)
            md_slides.append(slide_md)

            chunks.append(DocumentChunk(
                id=f"{doc_id}_s{slide_idx}",
                doc_id=doc_id,
                source_type=SourceType.PPTX,
                page_or_slide_num=slide_idx,
                text=slide_md,
                metadata={"slide_num": slide_idx, "filename": file_path.name}
            ))

        full_markdown = "\n\n---\n\n".join(md_slides)
        return ProcessedDocument(
            doc_id=doc_id,
            filename=file_path.name,
            source_type=SourceType.PPTX,
            title=file_path.stem.replace("_", " ").title(),
            markdown_content=full_markdown,
            chunks=chunks
        )

    def parse_text_or_md(self, file_path: Path) -> ProcessedDocument:
        """Parses Markdown or Text files using Semantic Header Chunking and parent-child tracking."""
        doc_id = self._clean_doc_id(file_path.stem)
        text = file_path.read_text(encoding="utf-8", errors="ignore")
        
        # Split semantically by Markdown headers (H1 to H4)
        raw_sections = re.split(r'(?m)(?=^#{1,4}\s+)', text)
        chunks: List[DocumentChunk] = []
        chunk_idx = 1

        for raw_sec in raw_sections:
            sec = raw_sec.strip()
            if not sec:
                continue

            # Identify section header if present
            header_match = re.match(r'^(#{1,4}\s+[^\n]+)', sec)
            header_title = header_match.group(1).strip() if header_match else "General"
            body = sec[len(header_match.group(0)):].strip() if header_match else sec

            paragraphs = [p.strip() for p in re.split(r'\n\s*\n', body) if p.strip()]
            if not paragraphs:
                chunks.append(DocumentChunk(
                    id=f"{doc_id}_c{chunk_idx}",
                    doc_id=doc_id,
                    source_type=SourceType.MARKDOWN,
                    page_or_slide_num=None,
                    text=sec,
                    metadata={
                        "chunk_index": chunk_idx,
                        "filename": file_path.name,
                        "section_title": header_title,
                        "parent_text": sec[:1500]
                    }
                ))
                chunk_idx += 1
                continue

            # Group paragraphs under current section logically
            curr_group = []
            curr_len = 0
            for p in paragraphs:
                curr_group.append(p)
                curr_len += len(p)
                if curr_len >= 500:
                    group_text = "\n\n".join(curr_group).strip()
                    chunk_text = f"{header_title}\n\n{group_text}" if header_title != "General" else group_text
                    chunks.append(DocumentChunk(
                        id=f"{doc_id}_c{chunk_idx}",
                        doc_id=doc_id,
                        source_type=SourceType.MARKDOWN,
                        page_or_slide_num=None,
                        text=chunk_text,
                        metadata={
                            "chunk_index": chunk_idx,
                            "filename": file_path.name,
                            "section_title": header_title,
                            "parent_text": sec[:1500]
                        }
                    ))
                    chunk_idx += 1
                    curr_group = []
                    curr_len = 0

            if curr_group:
                group_text = "\n\n".join(curr_group).strip()
                chunk_text = f"{header_title}\n\n{group_text}" if header_title != "General" else group_text
                chunks.append(DocumentChunk(
                    id=f"{doc_id}_c{chunk_idx}",
                    doc_id=doc_id,
                    source_type=SourceType.MARKDOWN,
                    page_or_slide_num=None,
                    text=chunk_text,
                    metadata={
                        "chunk_index": chunk_idx,
                        "filename": file_path.name,
                        "section_title": header_title,
                        "parent_text": sec[:1500]
                    }
                ))
                chunk_idx += 1

        return ProcessedDocument(
            doc_id=doc_id,
            filename=file_path.name,
            source_type=SourceType.MARKDOWN,
            title=file_path.stem.replace("_", " ").title(),
            markdown_content=text,
            chunks=chunks
        )

    def parse_and_export(self, file_path: Path) -> ProcessedDocument:
        """Parses document and exports normalized Markdown to data/markdown/ and data/vault/."""
        doc = self.parse_file(file_path)
        
        out_md = self.markdown_dir / f"{doc.doc_id}.md"
        out_md.write_text(doc.markdown_content, encoding="utf-8")

        vault_md = self.vault_dir / f"{doc.doc_id}.md"
        vault_md.write_text(doc.markdown_content, encoding="utf-8")

        logger.info(f"Ingested {file_path.name} -> {out_md} ({len(doc.chunks)} chunks)")
        return doc

    def _docx_table_to_markdown(self, table) -> str:
        """Converts docx.table.Table into GitHub Flavored Markdown table."""
        rows = []
        for r in table.rows:
            row_cells = [cell.text.replace("\n", " ").strip() for cell in r.cells]
            rows.append(row_cells)
        if not rows or not rows[0]:
            return ""
        
        header = "| " + " | ".join(rows[0]) + " |"
        sep = "| " + " | ".join(["---"] * len(rows[0])) + " |"
        body_lines = ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return "\n".join([header, sep] + body_lines)

    def _pptx_table_to_markdown(self, table) -> str:
        """Converts pptx Table into GitHub Flavored Markdown table."""
        rows = []
        for r in table.rows:
            row_cells = [cell.text.replace("\n", " ").strip() for cell in r.cells]
            rows.append(row_cells)
        if not rows or not rows[0]:
            return ""
        
        header = "| " + " | ".join(rows[0]) + " |"
        sep = "| " + " | ".join(["---"] * len(rows[0])) + " |"
        body_lines = ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return "\n".join([header, sep] + body_lines)

    def _clean_doc_id(self, name: str) -> str:
        clean = re.sub(r"[^a-zA-Z0-9_-]", "_", name)
        return clean.strip("_")
